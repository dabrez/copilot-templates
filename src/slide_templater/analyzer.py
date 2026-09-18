import json
import os
from pptx import Presentation
import spacy

def analyze_presentation(input_path, output_path):
    """
    Analyzes a .pptx presentation for potential variables using spaCy NER.
    
    :param input_path: Path to the source .pptx file.
    :param output_path: Path to save the generated JSON config file.
    """
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"Input file not found: {input_path}")

    # Load spaCy model
    try:
        nlp = spacy.load("en_core_web_trf")
    except OSError:
        raise OSError("spaCy model 'en_core_web_trf' not found. Please run 'python -m spacy download en_core_web_trf'")

    prs = Presentation(input_path)
    all_text = []

    # Extract all text from presentation
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame") and shape.text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    all_text.append(paragraph.text)
            
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text_frame:
                            for paragraph in cell.text_frame.paragraphs:
                                all_text.append(paragraph.text)

    # Process text with spaCy
    full_text = "\n".join(all_text)
    doc = nlp(full_text)

    # Extract entities
    # Labels of interest: DATE, PERSON, ORG, MONEY, GPE (Geopolitical Entity)
    labels_of_interest = {"DATE", "PERSON", "ORG", "MONEY", "GPE"}
    config = {}

    for ent in doc.ents:
        if ent.label_ in labels_of_interest:
            text = ent.text.strip()
            if text and text not in config:
                config[text] = f"[{ent.label_}]"

    # Save to JSON
    with open(output_path, 'w') as f:
        json.dump(config, f, indent=4)
