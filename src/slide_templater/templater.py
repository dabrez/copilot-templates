import re
from pptx import Presentation
import json
import os

def replace_text_in_presentation(input_path, config, output_path):
    """
    Replaces text in a .pptx presentation based on a configuration dictionary.
    
    :param input_path: Path to the source .pptx file.
    :param config: Dictionary where keys are text to find and values are replacements.
    :param output_path: Path to save the modified .pptx file.
    """
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"Input file not found: {input_path}")

    prs = Presentation(input_path)

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame") and shape.text_frame:
                _replace_text_in_text_frame(shape.text_frame, config)
            
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text_frame:
                            _replace_text_in_text_frame(cell.text_frame, config)

    prs.save(output_path)

def _apply(text, config):
    """
    Apply every replacement in a single pass.

    Sequential str.replace calls would feed each result back into the next
    key, so a config like {"Acme Corp": "Globex", "Globex": "WRONG"} would
    yield "WRONG". Matching all keys in one alternation means replacement
    output is never rescanned. Longest key first, so that an overlapping
    shorter key cannot claim part of a longer match.
    """
    if not config:
        return text
    pattern = "|".join(
        re.escape(k) for k in sorted(config, key=len, reverse=True) if k
    )
    if not pattern:
        return text
    return re.sub(pattern, lambda m: config[m.group(0)], text)


def _replace_text_in_paragraph(paragraph, config):
    """
    Replace text in a paragraph, handling cases where a match spans multiple runs.
    Collapses all runs into the first run when a cross-run match is found, preserving
    the first run's formatting as the best available approximation.
    """
    runs = paragraph.runs
    if not runs:
        return

    # Fast path: every match lies inside a single run, so per-run formatting
    # survives untouched. Counting rather than testing presence matters when a
    # key appears both intact in one run and split across a boundary -- the
    # intact copy would otherwise mask the split one.
    full_text = "".join(r.text for r in runs)
    spans_runs = any(
        full_text.count(k) > sum(r.text.count(k) for r in runs) for k in config
    )
    if not spans_runs:
        for run in runs:
            run.text = _apply(run.text, config)
        return

    # Slow path: a key straddles a run boundary, so the only way to match it
    # is on the joined text. That costs the other runs' formatting.
    runs[0].text = _apply(full_text, config)
    p_elem = paragraph._p
    for run in runs[1:]:
        p_elem.remove(run._r)


def _replace_text_in_text_frame(text_frame, config):
    for paragraph in text_frame.paragraphs:
        _replace_text_in_paragraph(paragraph, config)

def load_config(config_path):
    """Loads the JSON configuration file."""
    with open(config_path, 'r') as f:
        return json.load(f)
