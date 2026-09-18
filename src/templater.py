from pptx import Presentation
import json
import os

def replace_text_in_presentation(input_path, config, output_path):
    """
    Replaces text in a .pptx presentation based on a configuration dictionary.
    
    :param input_path: Path to the source .pptx file.
    :param config: Dictionary where keys are text to find and values are replacements.
    :param output_path: Path to save the modified .pptx file.
    """
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"Input file not found: {input_path}")

    prs = Presentation(input_path)

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame") and shape.text_frame:
                _replace_text_in_text_frame(shape.text_frame, config)
            
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text_frame:
                            _replace_text_in_text_frame(cell.text_frame, config)

    prs.save(output_path)

def _replace_text_in_paragraph(paragraph, config):
    """
    Replace text in a paragraph, handling cases where a match spans multiple runs.
    Collapses all runs into the first run when a cross-run match is found, preserving
    the first run's formatting as the best available approximation.
    """
    # Fast path: single-run or no match across runs
    for run in paragraph.runs:
        for find_text, replace_text in config.items():
            if find_text in run.text:
                run.text = run.text.replace(find_text, replace_text)

    # Slow path: check if any key spans multiple runs
    full_text = "".join(r.text for r in paragraph.runs)
    needs_merge = any(k in full_text and not any(k in r.text for r in paragraph.runs)
                      for k in config)
    if not needs_merge:
        return

    # Apply replacements on the merged text, then put it all in the first run
    for find_text, replace_text in config.items():
        full_text = full_text.replace(find_text, replace_text)

    from pptx.oxml.ns import qn
    runs = paragraph.runs
    if not runs:
        return
    runs[0].text = full_text
    # Remove extra runs from the XML
    p_elem = paragraph._p
    for run in runs[1:]:
        p_elem.remove(run._r)


def _replace_text_in_text_frame(text_frame, config):
    for paragraph in text_frame.paragraphs:
        _replace_text_in_paragraph(paragraph, config)

def load_config(config_path):
    """Loads the JSON configuration file."""
    with open(config_path, 'r') as f:
        return json.load(f)
