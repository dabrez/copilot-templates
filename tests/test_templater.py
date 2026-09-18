import json

import pytest
from pptx import Presentation
from pptx.util import Inches

from slide_templater.templater import (
    _replace_text_in_paragraph,
    load_config,
    replace_text_in_presentation,
)


def paragraph_with_runs(*parts):
    """A paragraph whose text is split across the given runs."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    paragraph = box.text_frame.paragraphs[0]
    for part in parts:
        paragraph.add_run().text = part
    return paragraph


def slide_text(path):
    return [
        p.text
        for slide in Presentation(path).slides
        for shape in slide.shapes
        if shape.has_text_frame
        for p in shape.text_frame.paragraphs
        if p.text.strip()
    ]


class TestParagraphReplacement:
    def test_replaces_within_a_single_run(self):
        p = paragraph_with_runs("Acme Corp annual review")
        _replace_text_in_paragraph(p, {"Acme Corp": "[ORG]"})
        assert p.text == "[ORG] annual review"

    def test_replaces_across_runs(self):
        # PowerPoint splits runs on formatting changes, so a visible phrase
        # is often not contained in any single run.
        p = paragraph_with_runs("Acme", " Corp")
        _replace_text_in_paragraph(p, {"Acme Corp": "[ORG]"})
        assert p.text == "[ORG]"

    def test_cross_run_match_collapses_into_one_run(self):
        p = paragraph_with_runs("Acme", " Corp")
        _replace_text_in_paragraph(p, {"Acme Corp": "[ORG]"})
        assert [r.text for r in p.runs] == ["[ORG]"]

    def test_handles_both_split_and_intact_matches_together(self):
        p = paragraph_with_runs("Acme Corp and Acme", " Corp")
        _replace_text_in_paragraph(p, {"Acme Corp": "[ORG]"})
        assert p.text == "[ORG] and [ORG]"

    def test_leaves_text_alone_when_nothing_matches(self):
        p = paragraph_with_runs("Nothing to see here")
        _replace_text_in_paragraph(p, {"Acme Corp": "[ORG]"})
        assert p.text == "Nothing to see here"

    def test_empty_config_is_a_noop(self):
        p = paragraph_with_runs("Acme", " Corp")
        _replace_text_in_paragraph(p, {})
        assert p.text == "Acme Corp"

    def test_replacements_do_not_cascade(self):
        """A replacement's output must not be rewritten by a later key.

        Replacing "Acme Corp" with "Globex" should leave "Globex" alone,
        even though "Globex" is itself a key in the same config.
        """
        p = paragraph_with_runs("Acme", " Corp")
        _replace_text_in_paragraph(p, {"Acme Corp": "Globex", "Globex": "WRONG"})
        assert p.text == "Globex"

    def test_no_cascade_within_a_single_run(self):
        p = paragraph_with_runs("Acme Corp")
        _replace_text_in_paragraph(p, {"Acme Corp": "Globex", "Globex": "WRONG"})
        assert p.text == "Globex"


class TestPresentation:
    def test_roundtrip_through_the_sample_deck(self, tmp_path, repo_root):
        out = tmp_path / "out.pptx"
        config = load_config(repo_root / "examples" / "config.json")
        replace_text_in_presentation(
            repo_root / "examples" / "sample.pptx", config, out
        )

        text = " ".join(slide_text(out))
        assert "Global Solutions Inc." in text
        assert "Acme Corp" not in text

    def test_replaces_inside_table_cells(self, tmp_path, repo_root):
        out = tmp_path / "out.pptx"
        replace_text_in_presentation(
            repo_root / "examples" / "sample.pptx", {"Acme Corp": "[ORG]"}, out
        )

        cells = [
            cell.text
            for slide in Presentation(out).slides
            for shape in slide.shapes
            if shape.has_table
            for row in shape.table.rows
            for cell in row.cells
        ]
        assert not any("Acme Corp" in c for c in cells)

    def test_leaves_the_input_file_untouched(self, tmp_path, repo_root):
        source = repo_root / "examples" / "sample.pptx"
        before = source.read_bytes()
        replace_text_in_presentation(
            source, {"Acme Corp": "[ORG]"}, tmp_path / "out.pptx"
        )
        assert source.read_bytes() == before

    def test_missing_input_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            replace_text_in_presentation(
                tmp_path / "nope.pptx", {}, tmp_path / "out.pptx"
            )


class TestLoadConfig:
    def test_reads_a_json_mapping(self, tmp_path):
        path = tmp_path / "c.json"
        path.write_text(json.dumps({"a": "b"}))
        assert load_config(path) == {"a": "b"}

    def test_malformed_json_raises(self, tmp_path):
        path = tmp_path / "c.json"
        path.write_text("{not json")
        with pytest.raises(json.JSONDecodeError):
            load_config(path)
